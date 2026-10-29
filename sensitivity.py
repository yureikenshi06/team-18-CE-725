"""One-way (tornado) and two-way sensitivity on equity NPV. Results are changes vs the base case."""
from dataclasses import replace
from engine import Inputs, run

def _fix(p: Inputs) -> Inputs:
    """Keep timing rules consistent when construction or sales windows change."""
    return replace(p, horizon=max(36, p.sales_end + 1, p.construction_months + 6),
                   holdback_month=p.construction_months + 1,
                   tax_month=max(36, p.sales_end + 1, p.construction_months + 6))

def vary(p: Inputs, name: str, shock: float) -> Inputs:
    """shock = +0.10 means a 10% adverse-or-favourable change in the named driver."""
    if name == "Sale price":          return _fix(replace(p, sale_price=p.sale_price * (1 + shock)))
    if name == "Construction cost":   return _fix(replace(p, construction_cost=p.construction_cost * (1 + shock)))
    if name == "Land cost":           return _fix(replace(p, land_price=p.land_price * (1 + shock)))
    if name == "Loan interest rate":  return _fix(replace(p, loan_rate=p.loan_rate * (1 + shock)))
    if name == "Construction period": return _fix(replace(p, construction_months=round(p.construction_months * (1 + shock))))
    if name == "Sales period":
        n = round((p.sales_end - p.sales_start + 1) * (1 + shock))
        return _fix(replace(p, sales_end=p.sales_start + n - 1))
    raise ValueError(name)

DRIVERS = ["Sale price", "Construction cost", "Land cost", "Loan interest rate",
           "Construction period", "Sales period"]

def tornado(p: Inputs, shock=0.10):
    base = run(p)[0]["equity_npv"]
    rows = []
    for d in DRIVERS:
        lo = run(vary(p, d, -shock))[0]["equity_npv"] - base
        hi = run(vary(p, d, +shock))[0]["equity_npv"] - base
        rows.append(dict(driver=d, minus=lo, plus=hi, swing=abs(hi - lo)))
    return sorted(rows, key=lambda r: r["swing"], reverse=True)

def two_way(p: Inputs, a: str, b: str, steps=(-0.10, -0.05, 0.0, 0.05, 0.10)):
    base = run(p)[0]["equity_npv"]
    return [[run(vary(vary(p, a, sa), b, sb))[0]["equity_npv"] - base for sb in steps] for sa in steps]

def plot_tornado(rows, path, shock=0.10):
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    rows = list(reversed(rows))
    fig, ax = plt.subplots(figsize=(9, 5))
    for i, r in enumerate(rows):
        ax.barh(i, r["minus"], color="#C55A11", height=0.55)
        ax.barh(i, r["plus"], color="#1F3864", height=0.55)
    ax.set_yticks(range(len(rows))); ax.set_yticklabels([r["driver"] for r in rows])
    ax.axvline(0, color="black", lw=0.8)
    ax.set_xlabel("Change in equity NPV vs base case (Rs crore)")
    ax.set_title(f"Tornado: each driver changed by -{int(shock*100)}% (orange) and +{int(shock*100)}% (blue)")
    for s in ("top", "right"): ax.spines[s].set_visible(False)
    fig.tight_layout(); fig.savefig(path, dpi=160); plt.close(fig)

def to_excel(p: Inputs, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Tornado"
    f = Font(name="Arial", size=10); h = Font(name="Arial", bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="1F3864")
    ws.append(["Driver", "-10% change in equity NPV (Rs cr)", "+10% change in equity NPV (Rs cr)", "Swing (Rs cr)"])
    for c in ws[1]: c.font, c.fill = h, fill
    for r in tornado(p): ws.append([r["driver"], r["minus"], r["plus"], r["swing"]])
    steps = (-0.10, -0.05, 0.0, 0.05, 0.10)
    for title, a, b in (("Price x Construction", "Sale price", "Construction cost"), ("Price x Sales period", "Sale price", "Sales period")):
        w = wb.create_sheet(title)
        w.append([f"Rows: {a}; columns: {b}. Values: change in equity NPV vs base (Rs cr)"])
        w.append([""] + [f"{int(s*100):+d}%" for s in steps])
        for s, row in zip(steps, two_way(p, a, b, steps)): w.append([f"{int(s*100):+d}%"] + row)
    for sh in wb:
        for row in sh.iter_rows():
            for c in row:
                if c.font != h: c.font = f
                if isinstance(c.value, float): c.number_format = "#,##0.00;(#,##0.00);-"
        sh.column_dimensions["A"].width = 24
        for col in "BCDEF": sh.column_dimensions[col].width = 30
    wb.save(path)

if __name__ == "__main__":
    p = Inputs()
    plot_tornado(tornado(p), "tornado.png"); to_excel(p, "sensitivity_tables.xlsx"); print("written")
